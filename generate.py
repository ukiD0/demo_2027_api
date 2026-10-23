# -*- coding: utf-8 -*-
"""
Генератор нового варианта КИМ Модуля 7 (09.02.07-5-2027).

Создаёт три файла:
  1) dataset_2026.xlsx       - «сырой» журнал транзакций с аномалиями (студентам);
  2) clean_dataset_2026.xlsx - эталон очищенных данных (результат предобработки, 7.1);
  3) ready_database.xlsx     - эталон: очистка + расчёты + график + ответы 7.2-7.11.

Требования: Python 3.8+, openpyxl  (pip install openpyxl)
Запуск:     python generate_kim7.py
"""

import re
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# 0. НАСТРОЙКИ
# ---------------------------------------------------------------------------
DATASET_FILE = "dataset_2026.xlsx"
CLEAN_FILE = "clean_dataset_2026.xlsx"
READY_FILE = "ready_database.xlsx"

MONTH_NAMES = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
               "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]

SEASONS = [("Зима", [1, 2]), ("Весна", [3, 4, 5]),
           ("Лето", [6, 7, 8]), ("Осень", [9, 10, 11])]

# Ожидаемые ответы — скрипт сверит с ними свой собственный расчёт
EXPECTED = {
    "7.2": "2370000",
    "7.3": "190909,09",
    "7.4": "700000,00",
    "7.5": "19,40",
    "7.6": "340000",
    "7.7": "3",
    "7.8": "Осень",
    "7.9": "7",
    "7.10": "Август",
    "7.11": "А",
}

# ---------------------------------------------------------------------------
# 1. «СЫРОЙ» ДАТАСЕТ — именно его получают студенты
# ---------------------------------------------------------------------------
# Формат: (transaction_id, date, product, category, quantity, price, shop, manager)
RAW_ROWS = [
    ("TXN-0003", 46029, "телевизор samsung", "Телевизоры", 2, 96000, "Яндекс.Маркет", "Иванов И.И."),
    ("TXN-0005", 46034, "Холодильник LG", "Телевизоры", "3 шт.", "84 000", "Эльдорадо", "смирнова е.в."),
    ("TXN-0011", 46041, "Чайник Kitfort ", "Чайники", 4, 4500, "ДНС ", "Петрова А.С."),
    ("TXN0014", "46048 ", "Микроволновая печь Midea", "Микроволновые печи", 4, 19500, "М.Видео", "Сидоров В.К."),
    ("TXN-0018", 46056, "ТЕЛЕВИЗОР SAMSUNG", "Телевизоры", 4, 124000, "Ситилинк", "ИВАНОВ И.И."),
    ("TXN-0023", 46062, "Холодильник LG", "Холодильники", " 4 ", "96 000 ₽", "Яндекс Маркет", "Смирнова Е.В."),
    ("TXN-0027", 46069, "Пылесос Dison", "Пылесосы", 2, 46000, "Эльдорадо", "Петрова А. С."),
    ("TXN-0031", "24.02.2026", "Стиральная машина Bosch", "Стиральные машины", 4, 52000, "ДНС", "Сидоров В.К."),
    ("TXN-0036", 46086, "Телевизор Samsung", "телевизоры", 2, 96000, "М.Видео", "Иванов И.И."),
    ("txn-0040", 46095, "холодильник Lg", "Утюги", 1, 84000, "ситилинк", "Смирнова Е.В."),
    ("TXN-0043", 46104, "Пылесос Dyson", "Пылесосы", "2 шт", "42 000", "Яндекс.Маркет", "Петрова А.С."),
    ("TXN-0047", 46114, "Пылесос  Dyson", "Пылесосы", 3, 124000, "Эльдорадо", "Сидоров  В.К."),
    ("TXN-0052", "46123", "Холодильник LG", "Холодильники", 2, 96000, "ДНС", "Иванов И.И."),
    ("TXN-0055", 46130, "утюг philips", "Утюги", 2, 5500, "М. Видео", "Смирнова Е.В."),
    ("TXN-0059", 46139, "Стиральная машина Bosch", "Стиральные машины", 3, "55 000,00", "Ситилинк", "иванов и.и"),
    (" TXN-0064 ", 46146, "Телевизор  Samsung", "Телевизоры", 5, 124000, "Яндекс.Маркет", "Петрова А.С."),
    ("TXN-0068", 46152, "Холодильник LG", "Холодильники", 3, 96000, "Эльдорадо", "Сидоров В.К."),
    ("TXN-0071", 46159, "Чайник Kitfort", "чайники", "4шт", 3500, "днс", "Петрова А.С."),
    ("TXN-0076", 46167, "Стиральная машина Bosch", "Стиральные машины", 4, "79 500", "М.Видео", "Иванов И.И."),
    ("TXN-0080", 46175, "телевизор samsung 4k", "Телевизоры", 2, 124000, "Ситилинк", "Смирнова Е.В."),
    ("TXN-0085", 46181, "Холодильник LG", "", 1, 84000, "Яндекс.Маркет", "Смирнова Е.В "),
    ("TXN-0089", 46188, "Пылесос Dyson", "Пылесосы", 2, 46000, "Эльдорадо", "Петрова А.С."),
    ("txn0092", "22.06.2026", "Микроволновая печь Mideа", "мелкaя бытовaя техникa", 2, 23000, "ДНС", "петрова а.с."),
    ("TXN-0097", 46209, "Телевизор Samsung", "Телевизоры", "1 шт.", 124000, "М.Видео", "Сидоров В.К."),
    ("TXN-0101", 46216, "стиральная машина bosch", "Стиральные машины", 3, 52000, "Ситилинк", "Иванов И.И."),
    ("TXN-0104", 46223, "Холодильник LG", "Холодильники", 1, 100000, "Яндекс.Маркет", "Сидоров В.К."),
    ("TXN-0108", 46237, "ТЕЛЕВИЗОР SAMSUNG", "Телевизоры", 5, "124 000", "Эльдорадо ", "Смирнова Е.В."),
    ("TXN-0113", 46243, "Холодильник LG", "Холодильники", "4 шт.", 96000, "ДНС", "Иванов И.И."),
    ("TXN-0117", "46250", "Чайник Kitfort", "Чайники", 4, 3500, "М.Видео", "Иванов И.И"),
    ("TXN-0120", 46258, "Стиральная Машина Bosch", "Стиральные машины", 4, "125 500", "Ситилинк", "Петрова А.С."),
    ("TXN0125", 46269, "Телевизор Samsung", "Телевизоры", 3, 124000, "Яндекс.Маркет ", "Сидоров В.К."),
    ("TXN-0129", 46276, "Холодильник LG", "Холодильники", 2, 96000, "Эльдорадо", "смирнова Е.В."),
    ("TXN-0132", 46283, "Утюг Philips ", "Пылесосы", 3, 5500, "ДНС", "Петрова А.С."),
    ("TXN-0136", "25.09.2026", "Стиральная машина Bosch", "Стиральные машины", 3, "46 500", "м.видео", "Сидоров В.К."),
    ("TXN-0141", 46297, "Телевизор Samsung", "Телевизоры", 2, 96000, "Ситилинк", "Иванов И.И."),
    ("TXN-0145", 46304, "Холодильник LG", " Телевизоры ", " 1 ", 84000, "Яндекс.Маркет", "Смирнова Е.В."),
    ("TXN-0148", 46311, "Чайник  Kitfort", "Чайники", 4, 3500, "Эльдорадо", "Петрова А.С."),
    ("TXN-0152", 46318, "Стиральная машина Bosch", "Стиральные машины", 4, "57 500 ₽", "ДНС", "Сидоров В.К."),
    (" TXN-0155", 46331, "телевизор Samsung", "Телевизоры", "1 шт.", 96000, "М.Видео", "Сидоров В.К."),
    ("TXN-0159", 46338, "Стиральная машина Bosch", "Стиральные машины", "3 шт.", 52000, "Ситилинк", "СИДОРОВ В.К."),
    ("TXN-0162", 46345, "Пылесос Dyson", "Пылесосы", 3, "46 000", "Яндекс Маркет", "Иванов И.И."),
    ("TXN-0166", 46359, "Телевизор Samsung ", "Телевизоры", 1, 124000, "Эльдорадо", "Петрова А.С."),
    ("TXN-0169", 46366, "Холодильник LG", "", 1, 84000, "ДНС", "Иванов И.И."),
    ("TXN-0173", "17.12.2026", "Пылесос Dison", "Пылесосы", 3, 44000, "М. Видео", "Иванов  И.И."),
]

# ---------------------------------------------------------------------------
# 2. ПРАВИЛА ОЧИСТКИ (те же действия, что должен выполнить студент)
# ---------------------------------------------------------------------------
HOMOGLYPHS = str.maketrans({
    "a": "а", "b": "в", "c": "с", "e": "е", "h": "н", "k": "к",
    "m": "м", "o": "о", "p": "р", "t": "т", "x": "х", "y": "у",
})


def fold(text):
    """Нижний регистр + замена латинских гомоглифов + сжатие пробелов."""
    s = str(text).strip().lower().translate(HOMOGLYPHS)
    return re.sub(r"[\s\u00a0]+", " ", s)


CATEGORY_BY_PRODUCT = {
    "Утюг Philips": "Утюги",
    "Микроволновая печь Midea": "Микроволновые печи",
    "Пылесос Dyson": "Пылесосы",
    "Чайник Kitfort": "Чайники",
    "Стиральная машина Bosch": "Стиральные машины",
    "Холодильник LG": "Холодильники",
    "Телевизор Samsung": "Телевизоры",
}

PRODUCT_RULES = [
    ("Утюг Philips", ("утюг", "philips")),
    ("Микроволновая печь Midea", ("микроволн", "midea")),
    ("Пылесос Dyson", ("пылесос", "dyson", "dison")),
    ("Чайник Kitfort", ("чайник", "kitfort")),
    ("Стиральная машина Bosch", ("стиральн", "bosch")),
    ("Холодильник LG", ("холодильник",)),
    ("Телевизор Samsung", ("телевизор", "samsung")),
]

SHOP_CANON = {
    "яндекс.маркет": "Яндекс.Маркет",
    "яндексмаркет": "Яндекс.Маркет",
    "эльдорадо": "Эльдорадо",
    "днс": "ДНС",
    "м.видео": "М.Видео",
    "ситилинк": "Ситилинк",
}

MANAGER_CANON = {
    re.sub(r"[\s.]", "", name.upper()): name
    for name in ("Иванов И.И.", "Смирнова Е.В.", "Петрова А.С.", "Сидоров В.К.")
}

EPOCH = datetime(1899, 12, 30)   # серийный номер 46023 = 01.01.2026


def clean_id(value):
    digits = re.sub(r"\D", "", str(value))
    return "TXN-{:04d}".format(int(digits))


def clean_date(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return EPOCH + timedelta(days=int(value))
    s = str(value).strip()
    if re.fullmatch(r"\d{4,5}", s):
        return EPOCH + timedelta(days=int(s))
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    raise ValueError("Не распознана дата: {!r}".format(value))


def clean_quantity(value):
    if isinstance(value, float) and value.is_integer():
        return int(value)
    digits = re.sub(r"\D", "", str(value))
    if not digits:
        raise ValueError("Не распознано количество: {!r}".format(value))
    return int(digits)


def clean_price(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(round(float(value)))
    s = str(value).strip().replace("\u00a0", " ")
    s = re.sub(r"(?i)(руб\.?|₽|rub)", "", s)
    s = s.replace(" ", "").replace(",", ".")
    return int(round(float(s)))


def clean_product(value):
    s = fold(value)
    for canonical, keys in PRODUCT_RULES:
        for key in keys:
            if fold(key) in s:
                return canonical
    raise ValueError("Не распознан товар: {!r}".format(value))


def clean_shop(value):
    key = re.sub(r"[\s\u00a0]+", "", str(value).strip().lower())
    if key not in SHOP_CANON:
        raise ValueError("Не распознан магазин: {!r}".format(value))
    return SHOP_CANON[key]


def clean_manager(value):
    key = re.sub(r"[\s.]", "", str(value).strip().upper())
    if key not in MANAGER_CANON:
        raise ValueError("Не распознан менеджер: {!r}".format(value))
    return MANAGER_CANON[key]


def build_clean():
    """Предобработка «сырого» файла -> очищенный набор строк."""
    rows = []
    for raw in RAW_ROWS:
        tid, date_v, product_v, _cat_v, qty_v, price_v, shop_v, mgr_v = raw
        d = clean_date(date_v)
        product = clean_product(product_v)
        quantity = clean_quantity(qty_v)
        price = clean_price(price_v)
        rows.append({
            "transaction_id": clean_id(tid),
            "date": d,
            "month": d.month,
            "product": product,
            "category": CATEGORY_BY_PRODUCT[product],
            "quantity": quantity,
            "price": price,
            "shop": clean_shop(shop_v),
            "manager": clean_manager(mgr_v),
            "revenue": quantity * price,
        })
    return rows


# ---------------------------------------------------------------------------
# 3. ОФОРМЛЕНИЕ
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def ru(x, nd=0):
    """Число в русском формате: 8 400 000 / 190 909,09"""
    return f"{x:,.{nd}f}".replace(",", " ").replace(".", ",")


def fmt2(x):
    """Ответ в формате тестирующей системы: 190909,09"""
    return f"{round(x, 2):.2f}".replace(".", ",")


# ---------------------------------------------------------------------------
# 4. ЗАПИСЬ ФАЙЛОВ
# ---------------------------------------------------------------------------
def write_dirty(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["transaction_id", "date", "product", "category",
               "quantity", "price", "shop", "manager"])
    for row in RAW_ROWS:
        ws.append(list(row))
    style_header(ws, 8)
    for i, w in enumerate([14, 12, 30, 26, 11, 12, 18, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def write_clean_sheet(ws, rows):
    ws.append(["transaction_id", "date", "product", "category", "quantity",
               "price", "shop", "manager", "Месяц", "Выручка за транзакцию"])
    for r in rows:
        ws.append([r["transaction_id"], r["date"], r["product"], r["category"],
                   r["quantity"], r["price"], r["shop"], r["manager"],
                   r["month"], r["revenue"]])
    style_header(ws, 10)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=2):
        row[0].number_format = "DD.MM.YYYY"
    for col in (6, 10):
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col, max_col=col):
            row[0].number_format = "#,##0"
    for i, w in enumerate([14, 12, 30, 24, 10, 12, 16, 16, 9, 20], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def write_ready(path, rows, monthly, total, avg_month, avg_txn, answers):
    wb = Workbook()

    # --- лист 1: очищенные данные ---
    ws1 = wb.active
    ws1.title = "clean_dataset_2026"
    write_clean_sheet(ws1, rows)

    # --- лист 2: расчёты + график ---
    ws2 = wb.create_sheet("Анализ")
    ws2["A1"] = "Месячная выручка за 2026 год"
    ws2["A1"].font = Font(bold=True, size=13)

    for i, h in enumerate(["Месяц", "Название", "Выручка, руб.", "Среднее значение, руб."], 1):
        ws2.cell(row=3, column=i, value=h)
    style_header(ws2, 4, row=3)

    for m in range(1, 13):
        ws2.cell(row=3 + m, column=1, value=m)
        ws2.cell(row=3 + m, column=2, value=MONTH_NAMES[m - 1])
        ws2.cell(row=3 + m, column=3, value=monthly[m]).number_format = "#,##0"
        ws2.cell(row=3 + m, column=4, value=avg_month).number_format = "#,##0"

    ws2.cell(row=16, column=1, value="Итого за год").font = Font(bold=True)
    ws2.cell(row=16, column=3, value=total).number_format = "#,##0"
    ws2.cell(row=17, column=1, value="Средняя месячная выручка").font = Font(bold=True)
    ws2.cell(row=17, column=3, value=avg_month).number_format = "#,##0.00"
    ws2.cell(row=18, column=1, value="Средняя выручка на транзакцию").font = Font(bold=True)
    ws2.cell(row=18, column=3, value=avg_txn).number_format = "#,##0.00"
    ws2.cell(row=18, column=4, value="транзакций: {}".format(len(rows)))

    # сезонная таблица
    ws2.cell(row=20, column=1, value="Сезонный анализ").font = Font(bold=True, size=12)
    for i, h in enumerate(["Сезон", "Месяцы", "Средняя выручка за месяц",
                           "Изменение к предыдущему сезону"], 1):
        ws2.cell(row=21, column=i, value=h)
    style_header(ws2, 4, row=21)

    season_avg = {}
    for idx, (name, months) in enumerate(SEASONS):
        season_avg[name] = sum(monthly[m] for m in months) / len(months)
        ws2.cell(row=22 + idx, column=1, value=name)
        ws2.cell(row=22 + idx, column=2,
                 value="{}-{}".format(min(months), max(months)))
        ws2.cell(row=22 + idx, column=3, value=season_avg[name]).number_format = "#,##0.00"
        if idx:
            prev = SEASONS[idx - 1][0]
            delta = season_avg[name] - season_avg[prev]
            ws2.cell(row=22 + idx, column=4, value=delta).number_format = "#,##0.00"

    # график: линия выручки + горизонтальная линия среднего
    chart = LineChart()
    chart.title = "Динамика выручки по месяцам, 2026"
    chart.style = 2
    chart.height = 9.5
    chart.width = 21
    chart.y_axis.title = "Выручка, руб."
    chart.x_axis.title = "Месяц"
    chart.y_axis.numFmt = "#,##0"

    data = Reference(ws2, min_col=3, min_row=3, max_row=15)
    cats = Reference(ws2, min_col=2, min_row=4, max_row=15)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.series[0].marker = Marker(symbol="circle", size=7)

    avg_ref = Reference(ws2, min_col=4, min_row=4, max_row=15)
    avg_series = Series(avg_ref, title="Среднее значение")
    avg_series.graphicalProperties.line.dashStyle = "dash"
    chart.series.append(avg_series)

    ws2.add_chart(chart, "F3")
    for i, w in enumerate([30, 28, 26, 30], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # --- лист 3: ответы ---
    ws3 = wb.create_sheet("Ответы 7.2-7.11")
    ws3["A1"] = "Эталонные ответы (формат ввода — как в тестирующей системе)"
    ws3["A1"].font = Font(bold=True, size=13)
    for i, h in enumerate(["№ задания", "Ответ", "Значение", "Как получено"], 1):
        ws3.cell(row=3, column=i, value=h)
    style_header(ws3, 4, row=3)

    for idx, (num, answer, value, comment) in enumerate(answers):
        ws3.cell(row=4 + idx, column=1, value=num)
        ws3.cell(row=4 + idx, column=2, value=answer).font = Font(bold=True)
        ws3.cell(row=4 + idx, column=3, value=value)
        ws3.cell(row=4 + idx, column=4, value=comment)

    r = 4 + len(answers) + 1
    ws3.cell(row=r, column=1, value="Контрольная сумма (7.1):").font = Font(bold=True)
    ws3.cell(row=r, column=2,
             value="выручка за год = {} руб.".format(ru(total)))
    ws3.cell(row=r + 1, column=1, value="Строк в очищенном файле:").font = Font(bold=True)
    ws3.cell(row=r + 1, column=2, value=len(rows))
    for i, w in enumerate([14, 18, 22, 70], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # --- лист 4: справка по очистке ---
    ws4 = wb.create_sheet("Справка по очистке")
    ws4["A1"] = "Аномалии исходного файла и правила приведения"
    ws4["A1"].font = Font(bold=True, size=13)
    notes = [
        "ID: 'TXN0014' (без дефиса), 'txn-0040', ' TXN-0064 ' (пробелы) -> TXN-00XX",
        "Дата: числа-серийники Excel, текст '46048', '46123'; даты 'ДД.ММ.ГГГГ' -> дата",
        "Товар: регистр, двойные/висячие пробелы, опечатки 'Dison', 'Mideа' (кириллица), 'samsung 4k'",
        "Категория: не соответствует товару, разный регистр, латиница, пустые значения -> восстановить по товару",
        "Количество: '3 шт.', '4шт', ' 1 ' -> целое число",
        "Цена: '84 000', '96 000 ₽', '55 000,00' -> целое число",
        "Магазин: регистр, пробелы, 'Яндекс Маркет' -> 'Яндекс.Маркет', 'М. Видео' -> 'М.Видео'",
        "Менеджер: регистр, пробелы, 'Петрова А. С.' -> 'Петрова А.С.'",
        "Полных дубликатов записей в файле нет — строки не удаляются.",
    ]
    for i, note in enumerate(notes, start=3):
        ws4.cell(row=i, column=1, value=note)
    ws4.column_dimensions["A"].width = 110

    wb.save(path)


# ---------------------------------------------------------------------------
# 5. ГЛАВНЫЙ СЦЕНАРИЙ
# ---------------------------------------------------------------------------
def main():
    rows = build_clean()

    ids = [r["transaction_id"] for r in rows]
    if len(ids) != len(set(ids)):
        print("ВНИМАНИЕ: обнаружены повторяющиеся transaction_id!")

    monthly = {m: 0 for m in range(1, 13)}
    for r in rows:
        monthly[r["month"]] += r["revenue"]

    total = sum(monthly.values())
    n = len(rows)
    avg_month = total / 12
    avg_txn = total / n

    summer = monthly[6] + monthly[7] + monthly[8]
    autumn = monthly[9] + monthly[10] + monthly[11]
    share_autumn = autumn / total * 100

    min_h2_month = min(range(7, 13), key=lambda m: monthly[m])
    min_h2 = monthly[min_h2_month]

    peaks = [m for m in range(2, 12)
             if monthly[m] > monthly[m - 1] and monthly[m] > monthly[m + 1]]
    below_avg = [m for m in range(1, 13) if monthly[m] < avg_month]

    diffs = {m: monthly[m] - monthly[m - 1] for m in range(2, 13)}
    growth_month = max(diffs, key=lambda m: diffs[m])

    season_avg = {name: sum(monthly[m] for m in ms) / len(ms)
                  for name, ms in SEASONS}
    season_change = {}
    for idx, (name, _) in enumerate(SEASONS):
        if idx:
            season_change[name] = season_avg[name] - season_avg[SEASONS[idx - 1][0]]
    drop_season = min(season_change, key=lambda k: season_change[k])

    s3, s4, s5 = monthly[3], monthly[4], monthly[5]
    if s3 < s4 < s5:
        ans_711 = "А"
    elif s3 > s4 > s5:
        ans_711 = "Б"
    else:
        ans_711 = "Г"

    computed = {
        "7.2": str(int(round(summer))),
        "7.3": fmt2(avg_txn),
        "7.4": fmt2(avg_month),
        "7.5": fmt2(share_autumn),
        "7.6": str(int(round(min_h2))),
        "7.7": str(len(peaks)),
        "7.8": drop_season,
        "7.9": str(len(below_avg)),
        "7.10": MONTH_NAMES[growth_month - 1],
        "7.11": ans_711,
    }

    blurbs = {
        "7.2": "Сумма выручки июнь-август: {} + {} + {}".format(
            ru(monthly[6]), ru(monthly[7]), ru(monthly[8])),
        "7.3": "Годовая выручка / число транзакций: {} / {}".format(ru(total), n),
        "7.4": "Годовая выручка / 12: {} / 12".format(ru(total)),
        "7.5": "Осень (сен-ноя) {} / год {} x 100%".format(ru(autumn), ru(total)),
        "7.6": "Минимум июля-декабря: {} ({})".format(
            ru(min_h2), MONTH_NAMES[min_h2_month - 1]),
        "7.7": "Локальные максимумы: {}".format(
            ", ".join(MONTH_NAMES[m - 1] for m in peaks)),
        "7.8": "Наибольшее снижение в сезоне '{}': {}".format(
            drop_season, ru(season_change[drop_season], 2)),
        "7.9": "Ниже средней {}: {}".format(
            ru(avg_month, 2), ", ".join(MONTH_NAMES[m - 1] for m in below_avg)),
        "7.10": "Максимальный прирост: {} (+{})".format(
            MONTH_NAMES[growth_month - 1], ru(diffs[growth_month])),
        "7.11": "Март {} < Апрель {} < Май {} - стабильный рост".format(
            ru(s3), ru(s4), ru(s5)),
    }

    answers = [(num, computed[num], computed[num], blurbs[num]) for num in EXPECTED]

    write_dirty(DATASET_FILE)
    wb_clean = Workbook()
    write_clean_sheet(wb_clean.active, rows)
    wb_clean.active.title = "clean_dataset_2026"
    wb_clean.save(CLEAN_FILE)
    write_ready(READY_FILE, rows, monthly, total, avg_month, avg_txn, answers)

    print("=" * 70)
    print("Созданы файлы:")
    print("  {:<26} - выдать студентам".format(DATASET_FILE))
    print("  {:<26} - эталон очистки (7.1)".format(CLEAN_FILE))
    print("  {:<26} - расчёты, график и ответы (7.2-7.11)".format(READY_FILE))
    print("=" * 70)
    print("Транзакций: {}".format(n))
    print("Выручка за год (контрольная сумма 7.1): {} руб.".format(ru(total)))
    print("Средняя месячная выручка: {}".format(ru(avg_month, 2)))
    print()
    print("Ответы:")
    for num in EXPECTED:
        print("  {:<6} {:<12} {}".format(num, computed[num], blurbs[num]))
    print()
    print("Сверка с ожидаемыми значениями:")
    ok = True
    for num, expected in EXPECTED.items():
        good = computed[num] == expected
        ok = ok and good
        print("  {:<6} {:>12}  {}".format(num, computed[num], "OK" if good else
                                          "РАСХОЖДЕНИЕ (ожидалось {})".format(expected)))
    print()
    print("Итог:", "всё сходится" if ok else "ПРОВЕРЬТЕ данные в RAW_ROWS")


if __name__ == "__main__":
    main()